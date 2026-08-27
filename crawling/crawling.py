import os
import re
import time
from urllib.parse import quote, urljoin

import requests
import pandas as pd
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage

from selenium import webdriver
from selenium.webdriver import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

BASE_DIR = os.getcwd()
IMAGE_DIR = os.path.join(BASE_DIR, "images")
EXCEL_FILE = os.path.join(BASE_DIR, "book_crawling_result.xlsx")
COLUMNS = ["검색어", "책제목", "저자", "가격", "출판사", "출판일", "이미지"]

SITE_URLS = {
    "yes24": "http://www.yes24.com/Main/default.aspx",
    "kyobo": "http://www.kyobobook.co.kr/",
    "aladin": "https://www.aladin.co.kr/home/welcome.aspx",
}

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0 Safari/537.36"
}


def create_folders():
    for site in ["yes24", "kyobo", "aladin"]:
        os.makedirs(os.path.join(IMAGE_DIR, site), exist_ok=True)


def create_driver():
    options = webdriver.ChromeOptions()
    options.add_argument("--start-maximized")
    options.add_argument("--lang=ko-KR")
    options.add_argument("--disable-notifications")
    driver = webdriver.Chrome(options=options)
    driver.set_page_load_timeout(30)
    return driver


def wait_body(driver, sec=10):
    WebDriverWait(driver, sec).until(
        EC.presence_of_element_located((By.TAG_NAME, "body"))
    )


def clean(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def clean_author(text):
    text = clean(text)
    for word in [
        "저자 모두보기", "저자(글)", "저자", "지음", "글", "옮김", "옮긴이",
        "번역", "감수", "원작", "더보기", "닫기", "(지은이)", "지은이", "외"
    ]:
        text = text.replace(word, " ")
    return clean(text).strip(" |/·,-")


def get_text(parent, selector):
    tag = parent.select_one(selector)
    return clean(tag.get_text(" ", strip=True)) if tag else ""


def get_lines(parent):
    return [clean(text) for text in parent.stripped_strings if clean(text)]


def extract_price(text):
    prices = re.findall(r"\d{1,3}(?:,\d{3})+\s*원|\d+\s*원", text or "")
    for price in prices:
        price = price.replace(" ", "")
        number = re.sub(r"[^0-9]", "", price)
        if number and int(number) > 0:
            return price
    return ""


def extract_date(text):
    patterns = [
        r"\d{4}년\s*\d{1,2}월\s*\d{1,2}일",
        r"\d{4}년\s*\d{1,2}월",
        r"\d{4}\.\s*\d{1,2}\.\s*\d{1,2}",
        r"\d{4}-\d{1,2}-\d{1,2}",
    ]
    for pattern in patterns:
        match = re.search(pattern, text or "")
        if match:
            return clean(match.group())
    return ""


def clean_img_url(img_url):
    if not img_url:
        return ""
    img_url = str(img_url).strip().strip("'\"").replace("&amp;", "&")
    if img_url.startswith("data:image"):
        return ""
    if "," in img_url:
        img_url = img_url.split(",")[0].strip()
    if " " in img_url and img_url.startswith(("http", "//", "/")):
        img_url = img_url.split()[0].strip()
    return img_url


def is_ready_image(img_url):
    lower = clean_img_url(img_url).lower()
    return any(word in lower for word in ["noimg", "no_img", "noimage", "ready", "loading", "blank", "sysimage"])


def get_img_from_box(box, base_url, site):
    attrs = ["src", "data-src", "data-original", "data-lazy", "data-kbbfn-img", "srcset"]
    domain = {
        "yes24": "image.yes24.com",
        "kyobo": "contents.kyobobook.co.kr",
        "aladin": "image.aladin.co.kr",
    }.get(site, "")

    for img in box.select("img"):
        for attr in attrs:
            img_url = clean_img_url(img.get(attr))
            if domain and domain in img_url and not is_ready_image(img_url):
                return urljoin(base_url, img_url)

    for img in box.select("img"):
        for attr in attrs:
            img_url = clean_img_url(img.get(attr))
            if not img_url:
                continue
            lower = img_url.lower()
            if any(word in lower for word in ["icon", "button", "logo", "blank", "loading", "banner", "noimg", "no_img", "noimage", "sysimage"]):
                continue
            return urljoin(base_url, img_url)
    return ""


def get_ext(img_url, content_type=""):
    content_type = (content_type or "").lower()
    if "png" in content_type:
        return ".png"
    if "webp" in content_type:
        return ".webp"
    if "gif" in content_type:
        return ".gif"

    path = img_url.split("?")[0].lower()
    for ext in [".jpg", ".jpeg", ".png", ".webp", ".gif"]:
        if path.endswith(ext):
            return ext
    return ".jpg"


def save_image(img_url, site, keyword, index, referer=""):
    try:
        img_url = clean_img_url(img_url)
        if not img_url:
            return ""

        if img_url.startswith("//"):
            img_url = "https:" + img_url
        elif img_url.startswith("/"):
            img_url = urljoin(SITE_URLS[site], img_url)

        headers = HEADERS.copy()
        headers["Referer"] = referer or SITE_URLS[site]
        response = requests.get(img_url, headers=headers, timeout=10)
        response.raise_for_status()

        if response.content[:30].lower().startswith((b"<!doctype", b"<html")):
            return ""

        ext = get_ext(img_url, response.headers.get("Content-Type", ""))
        safe_keyword = re.sub(r"[^가-힣a-zA-Z0-9]", "_", keyword)
        filename = f"{safe_keyword}_{index:04d}{ext}"
        save_path = os.path.join(IMAGE_DIR, site, filename)

        with open(save_path, "wb") as file:
            file.write(response.content)

        return os.path.relpath(save_path, BASE_DIR).replace("\\", "/")
    except Exception as e:
        print(f"[{site}] 이미지 다운로드 실패: {e}")
        return ""


def make_row(keyword, title, author, price, publisher, pub_date, img_url):
    return {
        "검색어": keyword,
        "책제목": title,
        "저자": author,
        "가격": price,
        "출판사": publisher,
        "출판일": pub_date,
        "이미지URL": img_url,
    }


def input_keyword(search_box, keyword):
    search_box.click()
    time.sleep(0.2)
    search_box.send_keys(Keys.CONTROL, "a")
    search_box.send_keys(Keys.BACKSPACE)
    time.sleep(0.2)
    search_box.send_keys(keyword)
    time.sleep(0.5)
    search_box.send_keys(Keys.ENTER)


def close_kyobo_popup(driver):
    for _ in range(5):
        try:
            buttons = driver.find_elements(
                By.XPATH,
                "//button[normalize-space()='확인'] | //a[normalize-space()='확인'] | //button[contains(@class, 'close')] | //a[contains(@class, 'close')]"
            )
            for button in buttons:
                if button.is_displayed():
                    driver.execute_script("arguments[0].click();", button)
                    time.sleep(0.4)
        except Exception:
            pass

        try:
            text = driver.find_element(By.TAG_NAME, "body").text
            if "오늘 하루" not in text and "그만 보시겠습니까" not in text:
                return
        except Exception:
            return

        try:
            ActionChains(driver).send_keys(Keys.ESCAPE).perform()
        except Exception:
            pass
        time.sleep(0.4)


def search_yes24(driver, keyword):
    driver.get(SITE_URLS["yes24"])
    wait_body(driver)
    time.sleep(1)
    search_box = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, "#query, input[name='query']"))
    )
    input_keyword(search_box, keyword)
    time.sleep(2)
    wait_body(driver)


def search_kyobo(driver, keyword):
    driver.get(SITE_URLS["kyobo"])
    wait_body(driver)
    time.sleep(2)
    close_kyobo_popup(driver)

    try:
        search_box = WebDriverWait(driver, 10).until(
            EC.element_to_be_clickable((By.CSS_SELECTOR, "input[name='keyword'], input#searchKeyword, input[placeholder*='검색']"))
        )
        input_keyword(search_box, keyword)
        time.sleep(2)
        wait_body(driver)
    except Exception:
        pass

    close_kyobo_popup(driver)
    if "search.kyobobook" not in driver.current_url:
        driver.get(f"https://search.kyobobook.co.kr/search?keyword={quote(keyword)}&gbCode=TOT&target=total")
        time.sleep(2)
        wait_body(driver)


def search_aladin(driver, keyword):
    driver.get(SITE_URLS["aladin"])
    wait_body(driver)
    time.sleep(1)

    try:
        search_box = WebDriverWait(driver, 10).until(
            EC.element_to_be_clickable((By.CSS_SELECTOR, "#SearchWord, input[name='SearchWord']"))
        )
        input_keyword(search_box, keyword)
        time.sleep(2)
        wait_body(driver)
    except Exception:
        pass

    if "wsearchresult" not in driver.current_url:
        driver.get(f"https://www.aladin.co.kr/search/wsearchresult.aspx?SearchTarget=Book&SearchWord={quote(keyword)}")
        time.sleep(2)
        wait_body(driver)


def parse_yes24(page_source, keyword):
    soup = BeautifulSoup(page_source, "html.parser")
    rows = []
    items = soup.select("li[data-goods-no]")[:24]

    for item in items:
        try:
            title_tag = item.select_one("a.gd_name, .goods_name a")
            title = clean(title_tag.get_text(" ", strip=True)) if title_tag else ""
            text = clean(item.get_text(" ", strip=True))

            author_tags = item.select(".info_auth a, .goods_auth a")
            author = ", ".join([clean_author(tag.get_text(" ", strip=True)) for tag in author_tags])
            if not author:
                author = clean_author(get_text(item, ".info_auth") or get_text(item, ".goods_auth"))

            publisher = get_text(item, ".info_pub") or get_text(item, ".goods_pub")
            pub_date = get_text(item, ".info_date") or extract_date(text)
            price = get_text(item, ".txt_num .yes_b") or get_text(item, ".yes_b") or extract_price(text)
            # YES24 목록 이미지에는 '이미지 준비중' placeholder가 섞일 수 있어서
            # 상품번호(data-goods-no)로 실제 표지 주소를 먼저 만든다.
            goods_no = item.get("data-goods-no", "")
            if not goods_no and title_tag:
                match = re.search(r"Goods/(\d+)", title_tag.get("href", ""))
                goods_no = match.group(1) if match else ""

            if goods_no:
                img_url = f"https://image.yes24.com/goods/{goods_no}/XL"
            else:
                img_url = get_img_from_box(item, SITE_URLS["yes24"], "yes24")

            rows.append(make_row(keyword, title, author, price, publisher, pub_date, img_url))
        except Exception as e:
            print(f"[yes24] 요소 탐색 실패: {e}")
    return rows


def parse_kyobo(page_source, keyword):
    soup = BeautifulSoup(page_source, "html.parser")
    rows = []
    items = soup.select("li.prod_item")[:20]

    for item in items:
        try:
            title_tag = item.select_one("a.prod_info .prod_name") or item.select_one("a.prod_info")
            link_tag = item.select_one("a.prod_info")
            title = clean(title_tag.get_text(" ", strip=True)) if title_tag else ""
            href = link_tag.get("href", "") if link_tag else ""
            text = clean(item.get_text(" ", strip=True))

            author = clean_author(get_text(item, ".prod_author") or get_text(item, ".author"))
            if not author:
                for line in get_lines(item):
                    if any(word in line for word in ["저자", "지음", "번역", "옮김"]):
                        author = clean_author(line.split("|")[0])
                        break

            pub_info = get_text(item, ".prod_publish") or get_text(item, ".publish")
            if not pub_info:
                for line in get_lines(item):
                    if "·" in line and extract_date(line):
                        pub_info = line
                        break

            pub_date = extract_date(pub_info)
            publisher = ""
            if pub_info and pub_date:
                before = pub_info.split(pub_date)[0]
                publisher = clean(before.rsplit("·", 1)[0] if "·" in before else before).strip(" |/·,-")

            price = get_text(item, ".price .val") or get_text(item, ".prod_price .val") or extract_price(text)
            img_url = get_img_from_box(item, SITE_URLS["kyobo"], "kyobo")

            if not img_url:
                match = re.search(r"97[89]\d{10}", str(item) + " " + href)
                if match:
                    img_url = f"https://contents.kyobobook.co.kr/sih/fit-in/458x0/pdt/{match.group()}.jpg"

            rows.append(make_row(keyword, title, author, price, publisher, pub_date, img_url))
        except Exception as e:
            print(f"[kyobo] 요소 탐색 실패: {e}")
    return rows


def clean_aladin_author_name(text, title=""):
    raw = clean(text)
    had_prefix = "[" in raw or " - " in raw

    if title:
        raw = raw.replace(title, " ")

    raw = re.sub(r"\[[^\]]+\]", " ", raw)
    raw = raw.replace("Choice", " ")

    if " - " in raw:
        raw = raw.split(" - ")[-1]

    raw = re.split(r"\(지은이\)|지은이|지음|저자|글|옮김|번역", raw)[0]
    raw = clean_author(raw)
    raw = clean(raw.replace(" ,", ",").replace(", ", ", "))

    bad_words = [
        "국내도서", "기본 개념", "세금", "전월세", "경매", "부의", "초격차",
        "레버리지", "투자 시스템", "BEST", "책자", "만드는", "입문", "자습서",
        "가이드", "미리보기", "장바구니", "보관함", "새창", "Sales Point"
    ]

    # 설명문 뒤에 저자명이 붙은 경우: 마지막의 '이름, 이름' 부분만 사용
    if had_prefix or len(raw) > 25 or any(word in raw for word in bad_words):
        match = re.search(r"([가-힣]{2,5}(?:\s*,\s*[가-힣]{2,5}){0,5})$", raw)
        if match:
            return clean(match.group(1))
        return ""

    return raw


def parse_aladin(page_source, keyword):
    soup = BeautifulSoup(page_source, "html.parser")
    rows = []
    items = soup.select("div.ss_book_box")[:25]

    for item in items:
        try:
            title_tag = item.select_one("a.bo3")
            title = clean(title_tag.get_text(" ", strip=True)) if title_tag else ""
            text = clean(item.get_text(" ", strip=True))

            author = ""
            publisher = ""
            pub_date = ""

            # 알라딘은 보통 한 li 안에 '저자 | 출판사 | 출판일'이 들어있다.
            for li in item.select("li"):
                li_text = clean(li.get_text(" ", strip=True))

                if not extract_date(li_text):
                    continue
                if "원" in li_text:
                    continue
                if any(word in li_text for word in ["장바구니", "보관함", "마이리스트", "미리보기", "새창"]):
                    continue

                pub_date = extract_date(li_text)

                links = [clean(a.get_text(" ", strip=True)) for a in li.select("a")]
                links = [x for x in links if x and x != title and "[" not in x]

                # a 태그가 있으면 마지막 a는 출판사, 그 앞은 저자로 사용
                if len(links) >= 2:
                    publisher = clean(links[-1])
                    names = []
                    for link in links[:-1]:
                        name = clean_aladin_author_name(link, title)
                        if name and name not in names:
                            names.append(name)
                    author = ", ".join(names)

                    if author and publisher:
                        break

                # a 태그로 부족하면 | 기준으로 분리
                parts = [clean(part) for part in li_text.split("|") if clean(part)]
                if len(parts) >= 3:
                    for i, part in enumerate(parts):
                        if extract_date(part) and i >= 2:
                            author = clean_aladin_author_name(parts[i - 2], title)
                            publisher = clean(parts[i - 1])
                            pub_date = extract_date(part) or part
                            break

                if author or publisher or pub_date:
                    break

            if not pub_date:
                pub_date = extract_date(text)

            price = get_text(item, ".ss_p2 b") or get_text(item, ".ss_p2") or extract_price(text)
            img_url = get_img_from_box(item, SITE_URLS["aladin"], "aladin")
            rows.append(make_row(keyword, title, author, price, publisher, pub_date, img_url))
        except Exception as e:
            print(f"[aladin] 요소 탐색 실패: {e}")
    return rows

def move_page(driver, site, keyword, page):
    try:
        driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
        time.sleep(1)
        page_btn = WebDriverWait(driver, 3).until(
            EC.element_to_be_clickable((By.XPATH, f"//a[normalize-space()='{page}']"))
        )
        driver.execute_script("arguments[0].click();", page_btn)
        time.sleep(2)
        wait_body(driver)
        return True
    except Exception:
        try:
            if site == "yes24":
                url = f"https://www.yes24.com/Product/Search?domain=ALL&query={quote(keyword)}&page={page}"
            elif site == "kyobo":
                url = f"https://search.kyobobook.co.kr/search?keyword={quote(keyword)}&gbCode=TOT&target=total&page={page}"
            else:
                url = f"https://www.aladin.co.kr/search/wsearchresult.aspx?SearchTarget=Book&SearchWord={quote(keyword)}&page={page}"
            driver.get(url)
            time.sleep(2)
            wait_body(driver)
            return True
        except Exception as e:
            print(f"[{site}] 페이지 이동 실패: {e}")
            return False


def crawl_site(driver, site, keyword, max_page):
    print(f"\n===== {site} 크롤링 시작 =====")
    start_time = time.perf_counter()
    all_rows = []

    search_funcs = {
        "yes24": search_yes24,
        "kyobo": search_kyobo,
        "aladin": search_aladin,
    }
    parse_funcs = {
        "yes24": parse_yes24,
        "kyobo": parse_kyobo,
        "aladin": parse_aladin,
    }

    try:
        search_funcs[site](driver, keyword)
    except Exception as e:
        print(f"[{site}] 검색 실패: {e}")
        return pd.DataFrame(columns=COLUMNS), 0

    for page in range(1, max_page + 1):
        print(f"[{site}] {page}페이지 수집 중...")
        try:
            wait_body(driver)
            time.sleep(0.5)
            rows = parse_funcs[site](driver.page_source, keyword)
        except Exception as e:
            print(f"[{site}] 요소 탐색 실패: {e}")
            rows = []

        if not rows:
            print(f"[{site}] 검색 결과가 없거나 부족하여 자동 종료합니다.")
            break

        for row in rows:
            img_path = save_image(
                row.get("이미지URL", ""),
                site,
                keyword,
                len(all_rows) + 1,
                driver.current_url,
            )
            all_rows.append({
                "검색어": row.get("검색어", ""),
                "책제목": row.get("책제목", ""),
                "저자": row.get("저자", ""),
                "가격": row.get("가격", ""),
                "출판사": row.get("출판사", ""),
                "출판일": row.get("출판일", ""),
                "이미지": img_path,
            })

        print(f"[{site}] {page}페이지 수집 완료: {len(rows)}건")

        if page < max_page:
            if not move_page(driver, site, keyword, page + 1):
                break

    elapsed = time.perf_counter() - start_time
    print(f"===== {site} 종료 / 총 {len(all_rows)}건 / {elapsed:.2f}초 =====")
    return pd.DataFrame(all_rows, columns=COLUMNS), elapsed


def insert_images_to_excel(excel_file):
    try:
        wb = load_workbook(excel_file)

        for sheet_name in ["yes24", "kyobo", "aladin"]:
            ws = wb[sheet_name]
            ws.column_dimensions["G"].width = 16

            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=7)
                img_path = cell.value

                if not img_path:
                    continue

                full_path = os.path.join(BASE_DIR, img_path)
                if not os.path.exists(full_path):
                    continue

                try:
                    img = ExcelImage(full_path)
                    img.width = 60
                    img.height = 85

                    # 이미지가 다음 행과 겹치지 않도록 행 높이를 이미지보다 넉넉하게 잡는다.
                    ws.row_dimensions[row].height = 105

                    # 이미지 경로 글자가 이미지 뒤에 보이지 않도록 셀 내용을 비운다.
                    cell.value = ""

                    # G열 셀 안에 이미지 삽입
                    ws.add_image(img, f"G{row}")

                except Exception:
                    continue

        wb.save(excel_file)

    except Exception as e:
        print(f"엑셀 이미지 삽입 실패: {e}")


def save_excel(site_dfs):
    try:
        with pd.ExcelWriter(EXCEL_FILE, engine="openpyxl") as writer:
            for sheet_name in ["yes24", "kyobo", "aladin"]:
                df = site_dfs.get(sheet_name, pd.DataFrame(columns=COLUMNS))
                df.to_excel(writer, sheet_name=sheet_name, index=False)

        insert_images_to_excel(EXCEL_FILE)
        print(f"\n엑셀 저장 완료: {EXCEL_FILE}")

    except PermissionError:
        new_file = os.path.join(BASE_DIR, f"book_crawling_result_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
        with pd.ExcelWriter(new_file, engine="openpyxl") as writer:
            for sheet_name in ["yes24", "kyobo", "aladin"]:
                df = site_dfs.get(sheet_name, pd.DataFrame(columns=COLUMNS))
                df.to_excel(writer, sheet_name=sheet_name, index=False)

        insert_images_to_excel(new_file)
        print("\n기존 엑셀 파일이 열려 있어서 다른 이름으로 저장했습니다.")
        print(new_file)

def crawl_start(keyword, yes24_pages, kyobo_pages, aladin_pages):
    total_start = time.perf_counter()
    create_folders()
    driver = create_driver()
    site_dfs = {}
    elapsed_times = {}

    try:
        for site, pages in [
            ("yes24", yes24_pages),
            ("kyobo", kyobo_pages),
            ("aladin", aladin_pages),
        ]:
            try:
                df, elapsed = crawl_site(driver, site, keyword, pages)
                site_dfs[site] = df
                elapsed_times[site] = elapsed
            except Exception as e:
                print(f"[{site}] 오류 발생. 다음 사이트로 넘어갑니다: {e}")
                site_dfs[site] = pd.DataFrame(columns=COLUMNS)
                elapsed_times[site] = 0
    finally:
        try:
            driver.quit()
        except Exception:
            pass

    save_excel(site_dfs)
    total_elapsed = time.perf_counter() - total_start

    print("\n===== 크롤링 시간 측정 =====")
    for site, elapsed in elapsed_times.items():
        print(f"{site}: {elapsed:.2f}초")
    print(f"전체: {total_elapsed:.2f}초")
    return site_dfs


if __name__ == "__main__":
    crawl_start("AI", 3, 6, 7)

